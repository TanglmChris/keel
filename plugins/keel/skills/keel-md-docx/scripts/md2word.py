#!/usr/bin/env python3
"""Convert Markdown and embedded SVG to DOCX, with SVG + PNG fallback images.

Usage: python md2word.py input.md -o output.docx
Dependencies: python-docx, lxml, Pillow; Node.js with marked and sharp.
"""
from __future__ import annotations

import argparse
import base64
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from urllib.parse import unquote, urlsplit

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Cm, Pt, RGBColor
from lxml import etree, html
from PIL import Image

HERE = Path(__file__).resolve().parent
SVG_NS = 'http://www.w3.org/2000/svg'
ASVG_NS = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'
SVG_PATTERN = re.compile(r'<svg\b[\s\S]*?</svg\s*>', re.IGNORECASE)


def node_call(payload: dict) -> object:
    node = os.environ.get('MD2WORD_NODE') or shutil.which('node')
    if not node:
        raise ValueError('Node.js not found; install Node.js or set MD2WORD_NODE.')
    result = subprocess.run([node, str(HERE / 'runtime.mjs')],
                            input=json.dumps(payload), text=True, capture_output=True,
                            timeout=90, check=False)
    if result.returncode:
        raise ValueError(result.stderr.strip() or 'Markdown/SVG helper failed')
    return json.loads(result.stdout)


class Converter:
    def __init__(self, source: Path, template: Path | None = None):
        self.source = source
        self.doc = Document(str(template)) if template else Document()
        if template:
            # Keep styles, headers and section settings; discard template body.
            body = self.doc._element.body
            for child in list(body):
                if child.tag != qn('w:sectPr'):
                    body.remove(child)
        else:
            section = self.doc.sections[0]
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            section.top_margin = section.bottom_margin = Cm(2)
            section.left_margin = section.right_margin = Cm(2)
            for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3',
                         'Heading 4', 'Heading 5', 'Heading 6', 'List Bullet', 'List Number']:
                style = self.doc.styles[name]
                style.font.name = 'Arial'
                style.font.color.rgb = RGBColor(0, 0, 0)
                rpr = style.element.get_or_add_rPr()
                fonts = rpr.find(qn('w:rFonts'))
                if fonts is None:
                    fonts = OxmlElement('w:rFonts'); rpr.append(fonts)
                fonts.set(qn('w:eastAsia'), 'Heiti SC')
            normal = self.doc.styles['Normal']
            normal.font.size = Pt(11)
            normal.paragraph_format.space_after = Pt(6)
            normal.paragraph_format.line_spacing = 1.15
        self.images = 0
        self.warnings: list[str] = []
        self.temp = tempfile.TemporaryDirectory(prefix='md2word-')

    @property
    def usable_size(self):
        s = self.doc.sections[-1]
        return (s.page_width - s.left_margin - s.right_margin,
                s.page_height - s.top_margin - s.bottom_margin - Cm(1))

    def local_image(self, href: str) -> tuple[bytes, str]:
        if href.startswith('data:'):
            header, payload = href.split(',', 1)
            data = base64.b64decode(payload, validate=True) if ';base64' in header else unquote(payload).encode()
            return data, 'svg' if 'image/svg+xml' in header else 'raster'
        split = urlsplit(href)
        if split.scheme or split.netloc:
            raise ValueError(f'Remote image is unsupported; save it locally first: {href}')
        target = (self.source.parent / unquote(split.path)).resolve()
        return target.read_bytes(), 'svg' if target.suffix.lower() == '.svg' else 'raster'

    def image(self, data: bytes, kind: str, parent=None, alt=''):
        svg = None
        if kind == 'svg':
            parser = etree.XMLParser(resolve_entities=False, no_network=True)
            root = etree.fromstring(data, parser)
            if etree.QName(root).localname != 'svg':
                raise ValueError('Expected an SVG root element')
            if root.getroottree().docinfo.doctype:
                raise ValueError('SVG DOCTYPE is unsupported')
            # A self-contained SVG makes the resulting DOCX portable.
            for element in root.iter():
                if not isinstance(element.tag, str):
                    continue
                tag = etree.QName(element).localname.lower()
                if tag in {'script', 'foreignObject'.lower()}:
                    raise ValueError(f'SVG {tag} is unsupported; use SVG text/path labels.')
                for key, value in element.attrib.items():
                    if etree.QName(key).localname.lower() in {'href', 'src'}:
                        if not value.startswith(('#', 'data:')):
                            raise ValueError('SVG external resource found; embed it as a data URI first.')
            text = data.decode('utf-8-sig')
            urls = re.findall(r'url\(\s*([^)]*)\)', text, re.IGNORECASE)
            if re.search(r'@import', text, re.IGNORECASE) or any(
                not value.strip().strip('\"\'').strip().startswith(('#', 'data:')) for value in urls
            ):
                raise ValueError('SVG external CSS/resource is unsupported')
            if not root.tag.startswith('{'):
                root.set('xmlns', SVG_NS)
            svg = etree.tostring(root, encoding='utf-8')
            target = Path(self.temp.name) / f'image-{self.images}.png'
            node_call({'action': 'svg', 'svg': svg.decode(), 'output': str(target)})
            data = target.read_bytes()
        with Image.open(io.BytesIO(data)) as picture:
            width, height = picture.size
        max_width, max_height = self.usable_size
        # Respect the source aspect ratio and avoid diagrams taller than a page.
        scale = min(max_width / width, max_height / height, 9525)
        p = parent if parent is not None else self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = p.add_run().add_picture(io.BytesIO(data), width=int(width * scale), height=int(height * scale))
        shape._inline.docPr.set('descr', alt or f'Figure {self.images + 1}')
        if svg is not None:
            existing = {str(part.partname) for part in self.doc.part.package.iter_parts()}
            index = 1
            while f'/word/media/vector{index}.svg' in existing:
                index += 1
            part = Part(PackURI(f'/word/media/vector{index}.svg'), 'image/svg+xml', svg, self.doc.part.package)
            rid = self.doc.part.relate_to(part, RT.IMAGE)
            blip = shape._inline.xpath('.//a:blip')[0]
            ext_list = OxmlElement('a:extLst')
            ext = OxmlElement('a:ext')
            ext.set('uri', '{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
            vector = etree.SubElement(ext, f'{{{ASVG_NS}}}svgBlip', nsmap={'asvg': ASVG_NS})
            vector.set(qn('r:embed'), rid)
            ext_list.append(ext); blip.append(ext_list)
        self.images += 1

    def inline(self, p, tokens: list[dict], bold=False, italic=False, strike=False, link=None):
        for token in tokens:
            kind = token['type']
            if kind in {'strong', 'em', 'del'}:
                self.inline(p, token.get('tokens', []), bold or kind == 'strong',
                            italic or kind == 'em', strike or kind == 'del', link)
            elif kind == 'link':
                self.inline(p, token.get('tokens', []), bold, italic, strike, token['href'])
            elif kind == 'image':
                data, image_kind = self.local_image(token['href'])
                self.image(data, image_kind, p, token.get('text', ''))
            elif kind == 'br':
                p.add_run().add_break()
            elif kind == 'html':
                text = token.get('text', '')
                if SVG_PATTERN.fullmatch(text.strip()):
                    self.image(text.encode(), 'svg', p)
                elif re.match(r'<br\s*/?>', text, re.IGNORECASE):
                    p.add_run().add_break()
                else:
                    self.warnings.append('Inline HTML styling is not preserved: ' + text[:60])
            else:
                text = token.get('text', token.get('raw', ''))
                if token.get('tokens'):
                    self.inline(p, token['tokens'], bold, italic, strike, link); continue
                run = p.add_run(html.fromstring(f'<span>{text}</span>').text_content() if '&' in text and '<' not in text else text)
                run.bold, run.italic, run.font.strike = bold, italic, strike
                if kind == 'codespan':
                    run.font.name = 'Consolas'; run.font.size = Pt(10)
                if link:
                    hyperlink = OxmlElement('w:hyperlink')
                    if link.startswith('#'):
                        # Heading bookmarks are not synthesized; keep the target visible.
                        self.warnings.append(f'Internal link target not resolved: {link}')
                    else:
                        hyperlink.set(qn('r:id'), self.doc.part.relate_to(link, RT.HYPERLINK, is_external=True))
                        p._p.remove(run._r); hyperlink.append(run._r); p._p.append(hyperlink)

    def raw_html(self, value: str):
        if re.search(r'<svg\b', value, re.IGNORECASE) and not SVG_PATTERN.search(value):
            raise ValueError('Unclosed SVG element')
        cursor = 0
        for match in SVG_PATTERN.finditer(value):
            self.html_fragment(value[cursor:match.start()])
            self.image(match.group().encode(), 'svg')
            cursor = match.end()
        self.html_fragment(value[cursor:])

    def html_fragment(self, value: str):
        if not value.strip():
            return
        try:
            root = html.fragment_fromstring(value, create_parent='div')
        except etree.ParserError:
            return
        for img in root.xpath('.//img'):
            data, kind = self.local_image(img.get('src', ''))
            self.image(data, kind, alt=img.get('alt', ''))
        text = root.text_content().strip()
        if text:
            self.doc.add_paragraph(text)
            self.warnings.append('HTML block preserved as plain text; use Markdown for formatting.')

    def blocks(self, tokens: list[dict], level=0):
        for token in tokens:
            kind = token['type']
            if kind in {'space', 'def'}:
                continue
            if kind == 'heading':
                p = self.doc.add_heading(level=token['depth'])
                self.inline(p, token.get('tokens', []))
            elif kind in {'paragraph', 'text'}:
                value = token.get('text', '')
                has_svg_html = any(child['type'] == 'html' and re.match(r'<svg\b', child.get('text', ''), re.IGNORECASE)
                                   for child in token.get('tokens', []))
                if has_svg_html:
                    # CommonMark treats a one-line SVG as inline HTML tokens.
                    # Read the original paragraph, not those individual tags.
                    self.raw_html(value)
                else:
                    p = self.doc.add_paragraph()
                    self.inline(p, token.get('tokens') or [{'type': 'text', 'text': value}])
            elif kind == 'code':
                if (token.get('lang') or '').strip().lower() == 'svg':
                    self.image(token['text'].encode(), 'svg')
                else:
                    p = self.doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(4)
                    run = p.add_run(token['text']); run.font.name = 'Consolas'; run.font.size = Pt(9)
            elif kind == 'html':
                self.raw_html(token['text'])
            elif kind == 'list':
                self.list_items(token, level)
            elif kind == 'blockquote':
                start = len(self.doc.paragraphs)
                self.blocks(token['tokens'], level)
                for p in self.doc.paragraphs[start:]:
                    p.paragraph_format.left_indent = Cm(.7)
            elif kind == 'table':
                self.table(token)
            elif kind == 'hr':
                self.doc.add_paragraph()
            else:
                raise ValueError(f'Unsupported Markdown block: {kind}')

    def list_items(self, token: dict, level: int):
        for index, item in enumerate(token['items']):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(.6 * (level + 1))
            p.paragraph_format.first_line_indent = Cm(-.4)
            # Explicit markers preserve arbitrary Markdown starting numbers across renderers.
            prefix = f"{int(token.get('start') or 1) + index}. " if token['ordered'] else '• '
            if item.get('task'):
                prefix = '[x] ' if item.get('checked') else '[ ] '
            p.add_run(prefix)
            first = True
            for child in item['tokens']:
                if child['type'] in {'text', 'paragraph'}:
                    if not first:
                        p = self.doc.add_paragraph()
                        p.paragraph_format.left_indent = Cm(.6 * (level + 1))
                    self.inline(p, child.get('tokens') or [{'type': 'text', 'text': child['text']}])
                    first = False
                elif child['type'] == 'list':
                    self.list_items(child, level + 1)
                else:
                    self.blocks([child], level + 1)

    def table(self, token: dict):
        rows = [token['header']] + token['rows']
        table = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = 'Table Grid'
        for i, row in enumerate(rows):
            trpr = table.rows[i]._tr.get_or_add_trPr()
            if i == 0:
                repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
            for j, cell in enumerate(row):
                target = table.cell(i, j)
                target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                self.inline(target.paragraphs[0], cell['tokens'])
                props = target._tc.get_or_add_tcPr()
                borders = OxmlElement('w:tcBorders')
                for side in ['top', 'left', 'bottom', 'right']:
                    line = OxmlElement('w:' + side)
                    for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                        line.set(qn('w:' + key), val)
                    borders.append(line)
                props.append(borders)
                margins = OxmlElement('w:tcMar')
                for side in ['top', 'left', 'bottom', 'right']:
                    margin = OxmlElement('w:' + side); margin.set(qn('w:w'), '90'); margin.set(qn('w:type'), 'dxa'); margins.append(margin)
                props.append(margins)
                if i == 0:
                    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EEF5'); props.append(shade)
                    for run in target.paragraphs[0].runs:
                        run.bold = True
                alignment = token.get('align', [None] * len(row))[j]
                if alignment in {'center', 'right'}:
                    target.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if alignment == 'center' else WD_ALIGN_PARAGRAPH.RIGHT
        self.doc.add_paragraph()

    def convert(self, output: Path):
        try:
            tokens = node_call({'action': 'parse', 'markdown': self.source.read_text(encoding='utf-8-sig')})
            self.blocks(tokens)
            output.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.NamedTemporaryFile(dir=output.parent, suffix='.docx', delete=False) as file:
                scratch = Path(file.name)
            try:
                self.doc.save(scratch)
                os.replace(scratch, output)
            finally:
                scratch.unlink(missing_ok=True)
        finally:
            self.temp.cleanup()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('-o', '--output', type=Path)
    parser.add_argument('--template', type=Path, help='Optional DOCX style template; its body is discarded')
    args = parser.parse_args()
    source = args.input.expanduser().resolve()
    output = (args.output or source.with_suffix('.docx')).expanduser().resolve()
    if source == output:
        parser.error('Input and output must be different files')
    try:
        converter = Converter(source, args.template)
        converter.convert(output)
        print(f'Created {output} ({converter.images} images)')
        for message in dict.fromkeys(converter.warnings):
            print('Warning: ' + message, file=sys.stderr)
    except (ValueError, OSError, etree.XMLSyntaxError, subprocess.TimeoutExpired) as exc:
        print(f'Conversion failed: {exc}', file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
